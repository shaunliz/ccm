# -*- coding: utf-8 -*-
"""사용자 가이드 슬라이드를 만드는 데 쓰는 뼈대.

한 자리에 모아 두는 이유는, 슬라이드마다 여백과 글꼴을 따로 적으면 장이
늘어날수록 표기가 갈리기 때문이다.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image

# --- 색 ---------------------------------------------------------------------
INK        = RGBColor(0x1F, 0x29, 0x37)   # 본문
INK_SOFT   = RGBColor(0x6B, 0x72, 0x80)   # 보조 설명
PRIMARY    = RGBColor(0x1B, 0x4F, 0x8A)   # 제목/강조 (프로그램의 파랑 계열)
ACCENT     = RGBColor(0x2E, 0x86, 0xDE)
GREEN      = RGBColor(0x1F, 0x9D, 0x55)
ORANGE     = RGBColor(0xD9, 0x7706 // 256 % 256, 0x06)  # 아래에서 다시 잡는다
ORANGE     = RGBColor(0xD9, 0x77, 0x06)
RED        = RGBColor(0xDC, 0x2E, 0x2E)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BG_SOFT    = RGBColor(0xF4, 0xF6, 0xF9)
RULE       = RGBColor(0xD7, 0xDD, 0xE5)

FONT = "맑은 고딕"
MONO = "Consolas"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.62)
CONTENT_TOP = Inches(1.55)
CONTENT_BOTTOM = Inches(6.95)


def _set_face(run, name):
    """라틴과 한글 글꼴을 함께 지정한다. ea 를 빼면 한글이 다른 글꼴로 나온다."""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def style_run(run, size=14, bold=False, color=INK, face=FONT, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    _set_face(run, face)
    return run


def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def textbox(slide, x, y, w, h, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def para(tf, text, size=14, bold=False, color=INK, face=FONT, space_after=6,
         space_before=0, align=PP_ALIGN.LEFT, level=0, line=1.25, first=False):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    style_run(r, size=size, bold=bold, color=color, face=face)
    return p


def rich(tf, chunks, size=14, space_after=6, space_before=0, level=0,
         align=PP_ALIGN.LEFT, line=1.25, first=False):
    """[(글자, {bold, color, face, size}), ...] 한 문단."""
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    for text, opt in chunks:
        r = p.add_run()
        r.text = text
        style_run(r, size=opt.get("size", size), bold=opt.get("bold", False),
                  color=opt.get("color", INK), face=opt.get("face", FONT))
    return p


def page_frame(slide, section, title, sub=None, page=None):
    """장마다 같은 머리/꼬리."""
    # 왼쪽 세로 강조선
    rect(slide, Inches(0), Inches(0), Inches(0.16), SLIDE_H, fill=PRIMARY)
    # 구역 이름
    _, tf = textbox(slide, MARGIN, Inches(0.42), Inches(9.0), Inches(0.3))
    para(tf, section, size=11, bold=True, color=ACCENT, space_after=0, first=True)
    # 제목
    _, tf = textbox(slide, MARGIN, Inches(0.72), Inches(11.9), Inches(0.6))
    para(tf, title, size=27, bold=True, color=PRIMARY, space_after=0, first=True)
    y = Inches(1.42)
    if sub:
        _, tf = textbox(slide, MARGIN, Inches(1.30), Inches(11.9), Inches(0.34))
        para(tf, sub, size=13, color=INK_SOFT, space_after=0, first=True)
        y = Inches(1.78)
    # 제목 아래 가로선
    rect(slide, MARGIN, y - Inches(0.10), Inches(12.1), Emu(9525), fill=RULE)
    # 꼬리
    if page is not None:
        _, tf = textbox(slide, Inches(11.9), Inches(7.02), Inches(0.9), Inches(0.3))
        para(tf, str(page), size=10, color=INK_SOFT, align=PP_ALIGN.RIGHT,
             space_after=0, first=True)
    _, tf = textbox(slide, MARGIN, Inches(7.02), Inches(7.0), Inches(0.3))
    para(tf, "Claude Code 사용량 모니터 v1.0.0 사용자 가이드", size=10,
         color=INK_SOFT, space_after=0, first=True)
    return y


def fit_image(slide, path, box_x, box_y, box_w, box_h, border=True, shadow=True):
    """상자 안에 비율을 지키며 넣는다. 가운데로 맞춘다."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(box_w / iw, box_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    x = int(box_x + (box_w - w) / 2)
    y = int(box_y + (box_h - h) / 2)
    if border:
        rect(slide, x - Emu(9525 * 1), y - Emu(9525 * 1), w + Emu(9525 * 2),
             h + Emu(9525 * 2), fill=None, line=RULE, line_w=0.75)
    pic = slide.shapes.add_picture(path, Emu(x), Emu(y), Emu(w), Emu(h))
    return pic, (x, y, w, h)


def caption(slide, text, x, y, w, align=PP_ALIGN.CENTER):
    _, tf = textbox(slide, x, y, w, Inches(0.3))
    para(tf, text, size=11, color=INK_SOFT, align=align, space_after=0, first=True)


def bullet_block(slide, items, x, y, w, h, size=14, gap=9, marker="•",
                 marker_color=ACCENT):
    """• 로 시작하는 목록. (글자) 또는 (굵은글자, 설명) 를 받는다."""
    _, tf = textbox(slide, x, y, w, h)
    first = True
    for it in items:
        if isinstance(it, tuple):
            head, tail = it
            # 굵은 머리말과 설명 사이는 반드시 벌어져야 한다. 보통 공백은 줄
            # 끝에서 지워지므로 em 공백을 쓴다.
            head = head.rstrip() + " "
            rich(tf, [(f"{marker} ", {"color": marker_color, "bold": True}),
                      (head, {"bold": True}),
                      (tail, {"color": INK})],
                 size=size, space_after=gap, first=first)
        else:
            rich(tf, [(f"{marker} ", {"color": marker_color, "bold": True}),
                      (it, {})], size=size, space_after=gap, first=first)
        first = False
    return tf


def table(slide, x, y, w, headers, rows, col_w=None, size=12, head_size=12,
          row_h=Inches(0.38), head_h=Inches(0.42)):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    h = head_h + row_h * len(rows)
    shape = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    tbl = shape.table
    tbl.first_row = True
    if col_w:
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = cw
    tbl.rows[0].height = head_h
    for i in range(len(rows)):
        tbl.rows[i + 1].height = row_h

    for c, htxt in enumerate(headers):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.10)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.line_spacing = 1.0
        r = p.add_run(); r.text = htxt
        style_run(r, size=head_size, bold=True, color=WHITE)

    for ri, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(ri + 1, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ri % 2 == 0 else BG_SOFT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.12
            face = MONO if (isinstance(val, str) and val.startswith("`")) else FONT
            text = val[1:-1] if face == MONO and val.endswith("`") else val
            r = p.add_run(); r.text = text
            style_run(r, size=size, color=INK, face=face,
                      bold=(c == 0 and n_cols > 1))
    return shape


def code_block(slide, lines, x, y, w, size=12, pad=Inches(0.16), title=None):
    h = pad * 2 + Inches(0.23) * len(lines) + (Inches(0.28) if title else 0)
    rect(slide, x, y, w, h, fill=RGBColor(0x1B, 0x1F, 0x27))
    _, tf = textbox(slide, x + pad, y + pad, w - pad * 2, h - pad * 2, wrap=False)
    first = True
    if title:
        para(tf, title, size=10, color=RGBColor(0x8A, 0x93, 0xA5), face=FONT,
             space_after=4, first=True)
        first = False
    for ln in lines:
        para(tf, ln, size=size, color=RGBColor(0xE6, 0xEA, 0xF0), face=MONO,
             space_after=1, line=1.15, first=first)
        first = False
    return h


def note_bar(slide, text, x, y, w, kind="info"):
    color = {"info": ACCENT, "warn": ORANGE, "ok": GREEN, "stop": RED}[kind]
    tint = {"info": RGBColor(0xEC, 0xF3, 0xFC), "warn": RGBColor(0xFD, 0xF4, 0xE6),
            "ok": RGBColor(0xEC, 0xF7, 0xF0), "stop": RGBColor(0xFD, 0xED, 0xED)}[kind]
    # 줄 수를 가늠해 높이를 잡는다. 고정 높이로 두면 긴 글이 상자 밖으로 넘친다.
    # 12pt 한글 한 자가 약 0.167인치다.
    usable = (w - Inches(0.55)) / Inches(1.0)
    per_line = max(10, int(usable / 0.167))
    lines = max(1, -(-len(text) // per_line))
    h = Inches(0.30 + 0.23 * lines)
    rect(slide, x, y, w, h, fill=tint)
    rect(slide, x, y, Inches(0.05), h, fill=color)
    _, tf = textbox(slide, x + Inches(0.20), y + Inches(0.12), w - Inches(0.38),
                    h - Inches(0.2))
    para(tf, text, size=12, color=INK, space_after=0, line=1.2, first=True)
    return h
